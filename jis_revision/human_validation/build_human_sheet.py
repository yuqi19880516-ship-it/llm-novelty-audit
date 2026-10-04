#!/usr/bin/env python3
"""Build the blinded human-coding workbook (2026-10-03): 150 drug–disease pairs drawn by stratified random sampling
(seed 20261005) from the 300 machine-read pairs, oversampling pairs the machine coders labelled L1/L2.
Coders see only drug, disease and the records; no machine labels. Answer key kept in a separate file."""
import json, glob, random
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
A = {}; [A.update({x["key"]: x for x in json.load(open(f))}) for f in glob.glob("coding_materials_v2/coder_A/*.json")]
items = {}; [items.update({x["key"]: x for x in json.load(open(f))}) for f in glob.glob("coding_materials_v2/batches/*.json")]
rng = random.Random(20261005)
by = {"L0": [], "L1": [], "L2": []}
for k, v in A.items(): by[v["label"]].append(k)
for v in by.values(): v.sort(); rng.shuffle(v)
pick = by["L0"][:150 - len(by["L1"][:35]) - len(by["L2"][:35])] + by["L1"][:35] + by["L2"][:35]
rng.shuffle(pick)
wb = Workbook(); ins = wb.active; ins.title = "Instructions"
for line in open("coding_materials_v2/codebook_v2.md").read().split("\n"): ins.append([line])
ins.column_dimensions["A"].width = 140
ws = wb.create_sheet("Coding")
ws.append(["no", "drug", "disease", "records (PMID | year | title | abstract)", "label (L0/L1/L2)", "deciding PMID", "deciding sentence", "comment"])
for i, k in enumerate(pick, 1):
    it = items[k]
    rec = "\n\n".join(f"[{r['pmid']} | {r['year']}] {r['title']}\n{r['abstract']}" for r in it["records"])
    ws.append([i, it["drug"], it["disease"], rec[:32000], "", "", "", ""])
    for c in ws[ws.max_row]: c.alignment = Alignment(wrap_text=True, vertical="top")
for col, w in zip("ABCDEFGH", [5, 18, 26, 110, 10, 12, 40, 20]): ws.column_dimensions[col].width = w
for c in ws[1]: c.font = Font(bold=True)
wb.save("coding_materials_v2/human_validation/human_coding_sheet_150.xlsx")
json.dump([{"no": i, "key": k, "stratum_machine_coderA": A[k]["label"]} for i, k in enumerate(pick, 1)],
          open("coding_materials_v2/human_validation/answer_key_DO_NOT_SEND.json", "w"), ensure_ascii=False, indent=1)
print(len(pick), "items")
