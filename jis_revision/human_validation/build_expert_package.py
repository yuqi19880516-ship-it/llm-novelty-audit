#!/usr/bin/env python3
"""Build the package sent to the two outside experts (2026-10-03).
Copies the blinded workbooks (no machine labels), adds a Chinese instruction sheet, drop-down validation for labels,
the target-equivalence rule for the free search (primary construct), and a merged-results template for the author."""
import shutil, os
from openpyxl import load_workbook, Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Alignment, Font, PatternFill
HV = os.path.dirname(os.path.abspath(__file__)); OUT = os.path.join(HV, "expert_package")
os.makedirs(OUT, exist_ok=True)

ZH_150 = """第一部分：150 对药物–疾病的文献判读（中文说明；以英文 Instructions 表的 codebook 为准）

每一行是一个药物 A 和一个目标疾病 C，D 列给出最多 6 条 PubMed 记录（题目和完整摘要），均发表于 2026 年 6 月 5 日之前。
只根据 D 列给出的记录判断，不要另外检索，也不要使用任何 AI 工具（ChatGPT、Claude、DeepSeek、Kimi 等）。
每行在 E 列下拉框中选一个标签：

L0 — 已有治疗提议：至少一条记录提出、试验、推荐或报告 A（或其盐、剂型）用于治疗、预防或辅助治疗 C 或 C 的亚型。
     任何研究类型都算：临床试验、观察性研究、治疗性使用的病例报告、C 的模型中的临床前研究、把 A 列为 C 候选疗法的综述。
L1 — 已有关联但无治疗提议：没有记录达到 L0，但至少一条记录陈述了 A 与 C 的实质性关系：
     机制或药理关系（如 A 作用于记录中与 C 相关的靶点或通路）、流行病学关联、A 引起或加重 C 的不良反应、药物–疾病相互作用数据。
L2 — 无已有关联：没有记录陈述 A 与 C 的关系。包括：没有记录；A 和 C 只是偶然同时出现（如某 C 患者因其他病服用 A 的病例、无关条目罗列）；
     匹配是假的（同名词、无关缩写）。

规则：
1. 关于另一种疾病、只是与 C 共享某个词的记录不算。
2. 关于药物类别的记录，只有明确点名 A 才算。
3. L0 与 L1 拿不准时选 L1；L1 与 L2 拿不准时，只有记录断言了关系才选 L1，单纯同时提到选 L2。
4. 疾病范围是 C 本身及其亚型；比 C 更宽的上位类别不等于 C（照 codebook 字面判断）。
5. F 列填决定标签的 PMID，G 列粘贴决定标签的那句原文；L2 填 none。H 列可写备注。

请独立完成，不与另一位专家讨论，完成并交回之前不要看对方的表。交回后不再修改 E 列。"""

ZH_30 = """第二部分：30 对药物–疾病的自由检索（中文说明）

对每一对，在 PubMed 和 Google Scholar 中自行检索 2026 年 6 月 5 日之前发表的任何文献（可用药物别名、商品名、疾病同义词）。
D 列下拉选择：
L0 — 有任一来源提出、试验或报告该药用于治疗该疾病（或其亚型）；
L1 — 没有 L0，但有来源陈述其他实质性关系（机制、关联、不良反应、相互作用）；
L2 — 两者都没找到。
E 列填决定性来源的 PMID 或 DOI，F 列填所用检索词。

疾病范围：
- 一般按疾病本身及其亚型。只涉及上位类别的文献（如目标为"难治性抑郁"而文献只写"抑郁症"；目标为"耐药结核"而文献只写"结核"）不计入 D 列，
  但请把这类文献的 PMID/DOI 记在 H 列"上位类别文献"中。
- 例外（三个目标的上位类别视同目标疾病，可以计入 D 列）：
  胰腺导管腺癌 ← 胰腺癌、胰腺腺癌；特发性肺纤维化 ← 肺纤维化（含博来霉素等肺纤维化模型）；
  非酒精性脂肪性肝炎 ← 非酒精性脂肪肝病（NAFLD）、代谢相关脂肪性肝病（MASLD）。
不要使用任何 AI 工具检索或判断。请独立完成，不与另一位专家讨论。"""

EN_FS_EXTRA = ["", "Disease scope: the target disease and its subtypes. Sources only on a broader category (e.g. 'depression' for treatment-resistant depression,",
    "'tuberculosis' for drug-resistant tuberculosis) do not count for column D; record them in column H ('broader-category source').",
    "Exception (target-equivalent categories, which do count for column D): pancreatic ductal adenocarcinoma <- pancreatic cancer / pancreatic adenocarcinoma;",
    "idiopathic pulmonary fibrosis <- pulmonary fibrosis, including bleomycin and similar models; non-alcoholic steatohepatitis <- NAFLD / MASLD.",
    "Do not use any AI tool (ChatGPT, Claude, DeepSeek, etc.) to search or to judge. Cut-off date: 5 June 2026."]

def finish(path, sheet, col, zh, extra_en=None, add_h=None):
    wb = load_workbook(path)
    if extra_en:
        for line in extra_en: wb["Instructions"].append([line])
    else:
        wb["Instructions"].append([""]); wb["Instructions"].append(["Do not search beyond the records shown and do not use any AI tool. Work independently of the other expert."])
    z = wb.create_sheet("说明(中文)", 0); [z.append([l]) for l in zh.split("\n")]; z.column_dimensions["A"].width = 150
    for r in z.iter_rows():
        for c in r: c.alignment = Alignment(wrap_text=True)
    ws = wb[sheet]
    if add_h: ws["H1"] = add_h; ws["H1"].font = Font(bold=True); ws.column_dimensions["H"].width = 30
    dv = DataValidation(type="list", formula1='"L0,L1,L2"', allow_blank=True, showErrorMessage=True,
                        errorTitle="标签", error="只能填 L0、L1 或 L2"); ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{ws.max_row}")
    yellow = PatternFill("solid", fgColor="FFF2CC")
    for r in range(2, ws.max_row + 1): ws[f"{col}{r}"].fill = yellow
    ws.freeze_panes = "B2"; wb.save(path)

for who in ("专家A", "专家B"):
    p1 = os.path.join(OUT, f"{who}_第一部分_150对文献判读.xlsx"); shutil.copy(os.path.join(HV, "human_coding_sheet_150.xlsx"), p1)
    finish(p1, "Coding", "E", ZH_150)
    p2 = os.path.join(OUT, f"{who}_第二部分_30对自由检索.xlsx"); shutil.copy(os.path.join(HV, "free_search_sheet_30.xlsx"), p2)
    finish(p2, "FreeSearch", "D", ZH_30, EN_FS_EXTRA, "broader-category source (上位类别文献)")

# merged results template (kept by the author; filled after both experts have returned their workbooks)
src1 = load_workbook(os.path.join(HV, "human_coding_sheet_150.xlsx"))["Coding"]
src2 = load_workbook(os.path.join(HV, "free_search_sheet_30.xlsx"))["FreeSearch"]
wb = Workbook(); ins = wb.active; ins.title = "填写说明"
for l in ["合并结果表（由作者在两位专家都交回后填写）",
          "1. 把专家A、专家B各自交回的标签原样粘贴到 D、E 列。粘贴后这两列不再修改。",
          "2. F 列自动比较：两人一致时显示同一标签；不一致时显示“待裁决”。",
          "3. 仅对“待裁决”的行，两位专家讨论（或请第三位专家）后，把结论填在 G 列。",
          "4. 两部分都填完后，把本文件连同四个原始工作表一起交回。"]: ins.append([l])
ins.column_dimensions["A"].width = 100
for title, src, n in (("第一部分_150", src1, 150), ("第二部分_30", src2, 30)):
    ws = wb.create_sheet(title)
    ws.append(["no", "drug", "disease", "专家A标签 expert A", "专家B标签 expert B", "是否一致 (自动)", "裁决标签 adjudicated (仅不一致行)", "裁决说明 note"])
    for r in src.iter_rows(min_row=2, values_only=True):
        if r[0] is None: continue
        i = ws.max_row + 1
        ws.append([r[0], r[1], r[2], None, None, f'=IF(OR(D{i}="",E{i}=""),"",IF(D{i}=E{i},D{i},"待裁决"))', None, None])
    dv = DataValidation(type="list", formula1='"L0,L1,L2"', allow_blank=True); ws.add_data_validation(dv)
    dv.add(f"D2:E{ws.max_row}"); dv.add(f"G2:G{ws.max_row}")
    for c in ws[1]: c.font = Font(bold=True)
    for col, w in zip("ABCDEFGH", [5, 20, 30, 14, 14, 14, 22, 40]): ws.column_dimensions[col].width = w
    ws.freeze_panes = "D2"
wb.save(os.path.join(OUT, "合并结果表_作者填写.xlsx"))
print(sorted(os.listdir(OUT)))
